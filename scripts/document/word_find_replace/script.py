"""
Word Document Find & Replace
Performs bulk find-and-replace in Word documents using JSON-formatted corrections.
"""

import json
from docx import Document


def find_and_replace_in_word(doc_path: str, corrections: list) -> int:
    """
    Find and replace text in a Word document.
    
    Args:
        doc_path: Path to the Word document (.docx)
        corrections: List of dicts with 'OgSentence' and 'NewSentence' keys
    
    Returns:
        Number of replacements made
    """
    try:
        document = Document(doc_path)
    except Exception as e:
        print(f"Error opening document: {e}")
        return 0

    replacements_made = 0

    for entry in corrections:
        original = entry.get("OgSentence")
        replacement = entry.get("NewSentence")

        if not original or not replacement:
            print(f"Skipping invalid entry: {entry}")
            continue

        # Check paragraphs
        for paragraph in document.paragraphs:
            if original in paragraph.text:
                paragraph.text = paragraph.text.replace(original, replacement, 1)
                replacements_made += 1
                break

        # Check tables
        for table in document.tables:
            found = False
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if original in paragraph.text:
                            paragraph.text = paragraph.text.replace(original, replacement, 1)
                            replacements_made += 1
                            found = True
                            break
                    if found:
                        break
                if found:
                    break
            if found:
                break

    # Save to new file
    output_path = doc_path.replace(".docx", "_updated.docx")
    try:
        document.save(output_path)
        print(f"Saved: {output_path}")
        print(f"Replacements: {replacements_made}")
    except Exception as e:
        print(f"Error saving: {e}")

    return replacements_made


def load_corrections_from_json(json_path: str) -> list:
    """Load corrections from a JSON file."""
    with open(json_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def parse_corrections_string(json_string: str) -> list:
    """
    Parse a string containing multiple JSON objects.
    Handles both array format and concatenated objects.
    """
    json_string = json_string.strip()
    
    # Try parsing as array first
    if json_string.startswith('['):
        return json.loads(json_string)
    
    # Parse concatenated JSON objects
    corrections = []
    decoder = json.JSONDecoder()
    idx = 0
    
    while idx < len(json_string):
        try:
            obj, new_idx = decoder.raw_decode(json_string[idx:])
            corrections.append(obj)
            idx += new_idx
            while idx < len(json_string) and json_string[idx].isspace():
                idx += 1
        except json.JSONDecodeError:
            idx += 1
    
    return corrections


def main():
    """
    Example usage - modify these values for your use case.
    """
    # Option 1: Load from JSON file
    # corrections = load_corrections_from_json("corrections.json")
    
    # Option 2: Define inline
    corrections = [
        {
            "OgSentence": "This is the original sentence.",
            "NewSentence": "This is the corrected sentence."
        },
        {
            "OgSentence": "Another sentence to fix.",
            "NewSentence": "Another sentence, now fixed."
        }
    ]
    
    doc_path = "document.docx"  # Change this to your document path
    
    print(f"Document: {doc_path}")
    print(f"Corrections: {len(corrections)}")
    print()
    
    find_and_replace_in_word(doc_path, corrections)
    print("Done!")


if __name__ == "__main__":
    main()
